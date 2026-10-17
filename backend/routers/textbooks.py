from auth import get_current_user
from auth import require_admin
from config import ZIP_BATCH_MAX_SIZE, MAX_IMPORT_BYTES
from database import get_db
from datetime import datetime
from fastapi import APIRouter
from fastapi import Depends
from fastapi import HTTPException
from fastapi.responses import FileResponse, Response, StreamingResponse
from fastapi.responses import Response
from models import NoOrderConfirm
from models import Textbook
from models import TextbookConfig
from models import TextbookOrder
from models import User
from schemas import TextbookBatchOrderRequest
from schemas import TextbookConfigCreate
from schemas import TextbookConfigResponse
from schemas import TextbookCreate
from schemas import TextbookOrderResponse
from schemas import TextbookResponse
from services.files import build_zip_response, safe_upload_path, save_upload
from services.mailer import load_smtp_config, save_smtp_config, send_emails_batch
from services.mailer import send_emails_batch
from services.qq import _get_qq_config
from services.qq import _get_qq_config, _send_qq_msg
from services.qq import _send_qq_msg
from services.utils import HAS_OPENPYXL
from services.utils import openpyxl
from sqlalchemy import func
from sqlalchemy.orm import Session
from starlette.background import BackgroundTask
from urllib.parse import quote
import io

router = APIRouter(tags=["textbooks"])


def get_textbook_config(db: Session) -> dict:
    """获取教材订购配置（deadline等），返回字典"""
    cfg = db.query(TextbookConfig).first()
    if cfg:
        return {
            "deadline": cfg.deadline,
            "notice_id": cfg.notice_id,
            "contact_user_id": cfg.contact_user_id
        }
    return {"deadline": None, "notice_id": None, "contact_user_id": None}



def check_deadline(db: Session):
    """检查是否在截止时间内，超过截止时间抛出403"""
    cfg = get_textbook_config(db)
    if cfg["deadline"]:
        now = datetime.now()
        limit = cfg["deadline"]
        if not isinstance(limit, datetime):
            limit = datetime.fromisoformat(str(limit))
        if now > limit:
            raise HTTPException(status_code=403, detail="教材订购已截止")



@router.get("/api/textbooks/config", response_model=TextbookConfigResponse)
def get_config(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(get_current_user)):
    """获取教材订购配置"""
    cfg = db.query(TextbookConfig).first()
    if not cfg:
        cfg = TextbookConfig()
        db.add(cfg)
        db.commit()
        db.refresh(cfg)
    return cfg



@router.put("/api/textbooks/config", response_model=TextbookConfigResponse)
def update_config(data: TextbookConfigCreate, db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """管理员设置教材订购配置"""
    cfg = db.query(TextbookConfig).first()
    if not cfg:
        cfg = TextbookConfig()
        db.add(cfg)
    if data.deadline is not None:
        cfg.deadline = data.deadline
    if data.notice_id is not None:
        cfg.notice_id = data.notice_id
    if data.contact_user_id is not None:
        cfg.contact_user_id = data.contact_user_id
    db.commit()
    db.refresh(cfg)
    return cfg



@router.get("/api/textbooks", response_model=list[TextbookResponse])
def list_textbooks(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(get_current_user)):
    """列出所有教材"""
    return db.query(Textbook).order_by(Textbook.created_at.desc()).all()



@router.post("/api/textbooks", response_model=TextbookResponse)
def create_textbook(data: TextbookCreate, db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """管理员新增教材"""
    t = Textbook(**data.model_dump())
    db.add(t)
    db.commit()
    db.refresh(t)
    return t



@router.put("/api/textbooks/{textbook_id}", response_model=TextbookResponse)
def update_textbook(textbook_id: int, data: TextbookCreate, db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """管理员修改教材"""
    t = db.query(Textbook).filter(Textbook.id == textbook_id).first()
    if not t: raise HTTPException(status_code=404, detail="教材不存在")
    for k, v in data.model_dump().items():
        setattr(t, k, v)
    db.commit()
    db.refresh(t)
    return t



@router.get("/api/textbooks/my-orders")
def get_my_orders(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(get_current_user)):
    """学生查看自己的所有订购"""
    orders = db.query(TextbookOrder).filter(
        TextbookOrder.student_id == current_user.id
    ).order_by(TextbookOrder.created_at.desc()).all()
    items = []
    total_original = 0.0
    total_discounted = 0.0
    for o in orders:
        t = o.textbook
        total_original += t.original_price if t else 0
        total_discounted += t.discounted_price if t else 0
        items.append({
            "id": o.id,
            "textbook_id": t.id if t else 0,
            "course_name": t.course_name or "" if t else "",
            "textbook_name": t.name if t else "",
            "isbn": t.isbn or "" if t else "",
            "publisher": t.publisher or "" if t else "",
            "author": t.author or "" if t else "",
            "course_type": t.course_type if t else "",
            "original_price": t.original_price if t else 0,
            "discounted_price": t.discounted_price if t else 0,
            "created_at": o.created_at.strftime("%Y-%m-%d %H:%M:%S") if o.created_at else ""
        })
    return {
        "orders": items,
        "total_original": round(total_original, 2),
        "total_discounted": round(total_discounted, 2),
        "no_order_confirmed": db.query(NoOrderConfirm).filter(NoOrderConfirm.student_id == current_user.id).first() is not None
    }



@router.delete("/api/textbooks/my-orders/{order_id}")
def cancel_my_order(order_id: int, db: Session = Depends(get_db, scope="function"), current_user: User = Depends(get_current_user)):
    """学生取消自己的某条订购"""
    check_deadline(db)
    order = db.query(TextbookOrder).filter(
        TextbookOrder.id == order_id,
        TextbookOrder.student_id == current_user.id
    ).first()
    if not order:
        raise HTTPException(status_code=404, detail="订购记录不存在")
    db.delete(order)
    db.commit()
    return {"message": "已取消订购"}



@router.post("/api/textbooks/order-batch")
def order_textbooks(data: TextbookBatchOrderRequest, db: Session = Depends(get_db, scope="function"), current_user: User = Depends(get_current_user)):
    """学生批量订购教材（重新订购时使用）"""
    check_deadline(db)
    if not data.textbook_ids:
        raise HTTPException(status_code=400, detail="请选择至少一本教材")
    textbooks = db.query(Textbook).filter(Textbook.id.in_(data.textbook_ids)).all()
    if len(textbooks) != len(data.textbook_ids):
        raise HTTPException(status_code=400, detail="部分教材不存在")
    # 清除该学生之前的订单
    db.query(TextbookOrder).filter(TextbookOrder.student_id == current_user.id).delete()
    # 清除无需订购确认
    db.query(NoOrderConfirm).filter(NoOrderConfirm.student_id == current_user.id).delete()
    # 创建新订单
    for tid in data.textbook_ids:
        order = TextbookOrder(textbook_id=tid, student_id=current_user.id)
        db.add(order)
    db.commit()
    return {"message": "订购成功", "count": len(data.textbook_ids)}



@router.delete("/api/textbooks/{textbook_id}")
def delete_textbook(textbook_id: int, db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """管理员删除教材（同时删除关联订单）"""
    t = db.query(Textbook).filter(Textbook.id == textbook_id).first()
    if not t: raise HTTPException(status_code=404, detail="教材不存在")
    db.query(TextbookOrder).filter(TextbookOrder.textbook_id == textbook_id).delete()
    db.delete(t)
    db.commit()
    return {"message": "删除成功"}



@router.get("/api/textbooks/orders", response_model=list[TextbookOrderResponse])
def list_all_orders(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """管理员查看全部订购记录"""
    orders = db.query(TextbookOrder).order_by(TextbookOrder.created_at.desc()).all()
    return [TextbookOrderResponse(
        id=o.id, textbook_id=o.textbook_id,
        course_name=o.textbook.course_name if o.textbook else "",
        textbook_name=o.textbook.name if o.textbook else "",
        isbn=o.textbook.isbn if o.textbook else "",
        publisher=o.textbook.publisher if o.textbook else "",
        author=o.textbook.author if o.textbook else "",
        course_type=o.textbook.course_type if o.textbook else "",
        original_price=o.textbook.original_price if o.textbook else 0,
        discounted_price=o.textbook.discounted_price if o.textbook else 0,
        student_id=o.student_id,
        student_name=o.student.real_name if o.student else "",
        student_student_id=o.student.student_id if o.student else "",
        created_at=o.created_at
    ) for o in orders]



@router.get("/api/textbooks/unordered")
def get_unordered_students(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """获取未订购教材的学生名单（含已确认无需订购）"""
    all_students = db.query(User).filter(User.role == "student").order_by(User.sort_order, User.student_id).all()
    ordered_ids = set(r[0] for r in db.query(TextbookOrder.student_id).distinct().all())
    no_order_ids = set(r[0] for r in db.query(NoOrderConfirm.student_id).distinct().all())

    unordered = []
    confirmed = []
    for s in all_students:
        item = {"id": s.id, "student_id": s.student_id, "real_name": s.real_name, "email": s.email}
        if s.id in no_order_ids:
            confirmed.append(item)
        elif s.id not in ordered_ids:
            unordered.append(item)
    return {
        "unordered": unordered, "count": len(unordered),
        "confirmed": confirmed, "confirmed_count": len(confirmed),
        "total": len(all_students)
    }



@router.post("/api/textbooks/confirm-no-order")
def confirm_no_order(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(get_current_user)):
    """学生确认无需订购教材"""
    check_deadline(db)
    # 清除已有订单
    db.query(TextbookOrder).filter(TextbookOrder.student_id == current_user.id).delete()
    # 记录确认
    existing = db.query(NoOrderConfirm).filter(NoOrderConfirm.student_id == current_user.id).first()
    if not existing:
        db.add(NoOrderConfirm(student_id=current_user.id))
    db.commit()
    return {"message": "已确认无需订购教材"}



@router.delete("/api/textbooks/confirm-no-order")
def cancel_no_order(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(get_current_user)):
    """学生取消无需订购确认"""
    check_deadline(db)
    db.query(NoOrderConfirm).filter(NoOrderConfirm.student_id == current_user.id).delete()
    db.commit()
    return {"message": "已取消"}



@router.get("/api/textbooks/orders/grouped")
def list_orders_grouped(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """管理员查看订购记录（按学生分组）"""
    orders = db.query(TextbookOrder).order_by(TextbookOrder.student_id, TextbookOrder.created_at).all()
    grouped: dict[int, dict] = {}
    for o in orders:
        sid = o.student_id
        if sid not in grouped:
            grouped[sid] = {
                "student_id": o.student.id,
                "student_name": o.student.real_name if o.student else "",
                "student_student_id": o.student.student_id if o.student else "",
                "student_email": o.student.email if o.student else "",
                "items": [],
                "total_original": 0,
                "total_discounted": 0,
                "latest_time": ""
            }
        t = o.textbook
        grouped[sid]["items"].append({
            "id": o.id,
            "textbook_id": t.id if t else 0,
            "course_name": t.course_name or "" if t else "",
            "textbook_name": t.name if t else "",
            "isbn": t.isbn or "" if t else "",
            "author": t.author or "" if t else "",
            "course_type": t.course_type if t else "",
            "original_price": t.original_price if t else 0,
            "discounted_price": t.discounted_price if t else 0,
            "created_at": o.created_at.strftime("%Y-%m-%d %H:%M:%S") if o.created_at else ""
        })
        if t:
            grouped[sid]["total_original"] += t.original_price or 0
            grouped[sid]["total_discounted"] += t.discounted_price or 0
    result = list(grouped.values())
    result.sort(key=lambda x: x["student_student_id"])
    # Round totals
    for r in result:
        r["total_original"] = round(r["total_original"], 2)
        r["total_discounted"] = round(r["total_discounted"], 2)
    return result



@router.post("/api/textbooks/qq-remind-unordered")
def qq_remind_unordered(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """QQ私聊提醒所有未订购学生"""
    cfg = _get_qq_config(db)
    if not cfg: raise HTTPException(status_code=400, detail="QQ提醒未启用或未配置")
    all_students = db.query(User).filter(User.role == "student").all()
    ordered_ids = set(r[0] for r in db.query(TextbookOrder.student_id).distinct().all())
    no_order_ids = set(r[0] for r in db.query(NoOrderConfirm.student_id).distinct().all())
    unordered = [s for s in all_students if s.id not in ordered_ids and s.id not in no_order_ids]
    if not unordered:
        return {"message": "所有学生均已订购或已确认无需订购", "sent": 0, "failed": 0}

    tcfg = get_textbook_config(db)
    deadline_str = ""
    if tcfg["deadline"]:
        deadline_str = tcfg["deadline"].strftime("%m月%d日 %H:%M") if hasattr(tcfg["deadline"], "strftime") else str(tcfg["deadline"])[:16]

    sent, failed = 0, 0
    for student in unordered:
        if not student.qq_number or not student.qq_number.strip():
            failed += 1
            continue
        msg = (
            f"【数科2441班·教材订购提醒】\n"
            f"同学 {student.real_name}，你好！\n"
            f"你尚未完成本学期的教材订购"
        )
        if deadline_str:
            msg += f"，截止时间为 {deadline_str}"
        msg += "。\n请尽快登录班务系统完成选购。"
        try:
            result = _send_qq_msg(cfg.api_url, "send_private_msg", {
                "user_id": int(student.qq_number.strip()),
                "message": msg,
            })
            if result.get("status") == "ok" or result.get("retcode") == 0:
                sent += 1
            else:
                failed += 1
        except Exception:
            failed += 1
    return {"message": f"已发送 {sent} 条QQ提醒" + (f"，{failed} 人发送失败" if failed > 0 else "")}



@router.post("/api/textbooks/qq-remind-unordered-email")
def qq_email_remind_unordered(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """QQ邮件提醒所有未订购学生"""
    all_students = db.query(User).filter(User.role == "student").all()
    ordered_ids = set(r[0] for r in db.query(TextbookOrder.student_id).distinct().all())
    no_order_ids = set(r[0] for r in db.query(NoOrderConfirm.student_id).distinct().all())
    unordered = [s for s in all_students if s.id not in ordered_ids and s.id not in no_order_ids]

    tcfg = get_textbook_config(db)
    deadline_str = ""
    if tcfg["deadline"]:
        deadline_str = tcfg["deadline"].strftime("%m月%d日 %H:%M") if hasattr(tcfg["deadline"], "strftime") else str(tcfg["deadline"])[:16]

    emails = []
    for student in unordered:
        if student.email and "@" in student.email:
            emails.append(student.email)
    if not emails:
        raise HTTPException(status_code=400, detail="所有学生均已订购或已确认无需订购，或无可用邮箱")

    lines = [f"同学，你好！\n"]
    lines.append("你尚未完成本学期的教材订购。")
    if deadline_str:
        lines.append(f"截止时间为：{deadline_str}")
    lines.append("请尽快登录班务系统（数科2441智能管理平台）完成教材选购。\n")
    lines.append("登录地址：https://你的服务器地址")
    lines.append("\n---\n此邮件由智能班务系统自动发送，请勿回复。")

    result = send_emails_batch(emails, "【数科2441班】教材订购提醒", "\n".join(lines), db=db, send_type="textbook_remind")
    return {"message": "已发送" if result["success"] else "发送失败"}



def export_orders_excel(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """导出订购明细 Excel"""
    if not HAS_OPENPYXL:
        raise HTTPException(status_code=500, detail="服务器未安装 openpyxl")
    orders = db.query(TextbookOrder).order_by(TextbookOrder.created_at.desc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "教材订购明细"
    ws.append(["序号", "学号", "姓名", "课程名称", "教材名称", "ISBN", "出版社", "作者", "课程类型", "单价", "折后单价", "订购时间"])

    total_original = 0.0
    total_discounted = 0.0
    for idx, o in enumerate(orders, 1):
        t = o.textbook
        ws.append([
            idx,
            o.student.student_id if o.student else "",
            o.student.real_name if o.student else "",
            t.course_name if t else "",
            t.name if t else "",
            t.isbn if t else "",
            t.publisher if t else "",
            t.author if t else "",
            t.course_type if t else "",
            t.original_price if t else 0,
            t.discounted_price if t else 0,
            o.created_at.strftime("%Y-%m-%d %H:%M:%S") if o.created_at else ""
        ])
        total_original += t.original_price if t else 0
        total_discounted += t.discounted_price if t else 0

    ws.append([])
    ws.append(["", "", "", "", "", "", "", "", "合计", round(total_original, 2), round(total_discounted, 2), ""])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = quote("教材订购明细.xlsx")
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"})



@router.post("/api/textbooks/orders/send-emails")
def send_order_emails(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """管理员一键发送订购明细到各学生邮箱"""
    orders = db.query(TextbookOrder).order_by(TextbookOrder.student_id).all()
    student_orders: dict[int, dict] = {}
    for o in orders:
        sid = o.student_id
        if sid not in student_orders:
            student_orders[sid] = {
                "student": o.student,
                "items": []
            }
        t = o.textbook
        student_orders[sid]["items"].append({
            "course_name": t.course_name or "" if t else "",
            "textbook_name": t.name if t else "",
            "isbn": t.isbn or "" if t else "",
            "author": t.author or "" if t else "",
            "original_price": t.original_price if t else 0,
            "discounted_price": t.discounted_price if t else 0,
        })

    if not student_orders:
        raise HTTPException(status_code=400, detail="暂无订购记录")

    sent = 0
    failed = 0
    for sid, data in student_orders.items():
        student = data["student"]
        if not student or not student.email or "@" not in student.email:
            failed += 1
            continue
        items = data["items"]
        total_orig = sum(i["original_price"] for i in items)
        total_disc = sum(i["discounted_price"] for i in items)

        lines = [f"{student.real_name or student.username} 同学，你好！\n"]
        lines.append("以下是你的教材订购明细：\n")
        for idx, it in enumerate(items, 1):
            lines.append(f"  {idx}. [{it['course_name']}] {it['textbook_name']}")
            if it['isbn']:
                lines.append(f"     ISBN: {it['isbn']}")
            if it['author']:
                lines.append(f"     作者: {it['author']}")
            lines.append(f"     单价: ¥{it['original_price']}  折后价: ¥{it['discounted_price']}")
            lines.append("")
        lines.append(f"原价合计: ¥{round(total_orig, 2)}")
        lines.append(f"折后合计: ¥{round(total_disc, 2)}")
        lines.append(f"\n共 {len(items)} 本教材，请核对后按班级通知的方式完成付款。")

        try:
            result = send_emails_batch(
                [student.email],
                "教材订购明细确认",
                "\n".join(lines),
                db=db,
                send_type="textbook_order",
                related_id=sid
            )
            if result["success"] and result["sent"] > 0:
                sent += 1
            else:
                failed += 1
        except Exception:
            failed += 1

    return {"message": f"已发送 {sent} 封邮件" + (f"，{failed} 人发送失败" if failed > 0 else "")}



@router.delete("/api/textbooks/orders/{order_id}")
def admin_delete_order(order_id: int, db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """管理员删除某条订购记录"""
    order = db.query(TextbookOrder).filter(TextbookOrder.id == order_id).first()
    if not order: raise HTTPException(status_code=404, detail="订购记录不存在")
    db.delete(order)
    db.commit()
    return {"message": "删除成功"}



@router.get("/api/textbooks/stats")
def get_textbook_stats(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """获取每本教材的订购数量统计"""
    textbooks = db.query(Textbook).all()
    result = []
    for t in textbooks:
        count = db.query(TextbookOrder).filter(TextbookOrder.textbook_id == t.id).count()
        result.append({
            "id": t.id,
            "name": t.name,
            "course_name": t.course_name or "",
            "course_type": t.course_type,
            "order_count": count
        })
    return result



def _get_dorm_grouped_data(db: Session):
    """获取按宿舍分组的订购数据"""
    orders = db.query(TextbookOrder).join(Textbook).join(User, TextbookOrder.student_id == User.id).all()
    # 按宿舍分组
    dorms: dict[str, dict] = {}
    for o in orders:
        dorm_name = (o.student.dorm or "未分配宿舍").strip()
        if dorm_name not in dorms:
            dorms[dorm_name] = {"students": [], "book_counts": {}}
        sid = o.student_id
        if sid not in dorms[dorm_name]["students"]:
            dorms[dorm_name]["students"].append(sid)
        tid = o.textbook_id
        if tid not in dorms[dorm_name]["book_counts"]:
            dorms[dorm_name]["book_counts"][tid] = {"name": o.textbook.name, "course_name": o.textbook.course_name or "", "count": 0}
        dorms[dorm_name]["book_counts"][tid]["count"] += 1
    return dorms



def _get_student_dorm_data(db: Session):
    """获取每个学生+宿舍的订购明细"""
    orders = db.query(TextbookOrder).order_by(TextbookOrder.student_id).all()
    students: dict[int, dict] = {}
    for o in orders:
        sid = o.student_id
        if sid not in students:
            students[sid] = {
                "name": o.student.real_name if o.student else "",
                "student_id": o.student.student_id if o.student else "",
                "dorm": (o.student.dorm or "").strip() if o.student else "",
                "items": []
            }
        students[sid]["items"].append({
            "name": o.textbook.name if o.textbook else "",
            "course_name": o.textbook.course_name or "" if o.textbook else "",
            "discounted_price": o.textbook.discounted_price if o.textbook else 0,
            "original_price": o.textbook.original_price if o.textbook else 0,
        })
    return students



@router.get("/api/textbooks/orders/export-by-dorm-excel")
def export_orders_by_dorm_excel(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """按宿舍导出 Excel"""
    if not HAS_OPENPYXL:
        raise HTTPException(status_code=500, detail="服务器未安装 openpyxl")
    students = _get_student_dorm_data(db)
    dorms = _get_dorm_grouped_data(db)

    wb = openpyxl.Workbook()
    # Sheet 1: 宿舍汇总
    ws1 = wb.active
    ws1.title = "宿舍订购汇总"
    header_font = openpyxl.styles.Font(bold=True, size=11)
    header_fill = openpyxl.styles.PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_w = openpyxl.styles.Font(bold=True, size=11, color="FFFFFF")
    center = openpyxl.styles.Alignment(horizontal="center", vertical="center")

    ws1.append(["宿舍号", "教材名称", "课程名称", "订购数量"])
    for cell in ws1[1]: cell.font = header_font_w; cell.fill = header_fill; cell.alignment = center

    for dorm_name in sorted(dorms.keys()):
        data = dorms[dorm_name]
        row_start = ws1.max_row + 1
        for tid, info in data["book_counts"].items():
            ws1.append([dorm_name, info["name"], info["course_name"], info["count"]])
        # 合计行
        total = sum(info["count"] for info in data["book_counts"].values())
        ws1.append([dorm_name, "合计", "", total])
        ws1.cell(row=ws1.max_row, column=2).font = openpyxl.styles.Font(bold=True)

    # Sheet 2: 学生明细
    ws2 = wb.create_sheet("学生订购明细")
    ws2.append(["学号", "姓名", "宿舍", "教材名称", "课程名称", "折后单价", "订购数量", "折后小计"])
    for cell in ws2[1]: cell.font = header_font_w; cell.fill = header_fill; cell.alignment = center

    total_all = 0
    for sid in sorted(students.keys(), key=lambda x: students[x]["student_id"]):
        s = students[sid]
        total_student = sum(i["discounted_price"] for i in s["items"])
        total_all += total_student
        ws2.append([s["student_id"], s["name"], s["dorm"], "", "", "", len(s["items"]), round(total_student, 2)])
        last_row = ws2.max_row
        for c in range(1, 9): ws2.cell(row=last_row, column=c).font = openpyxl.styles.Font(bold=True, color="1F4E79")
        for item in s["items"]:
            ws2.append(["", "", "", item["name"], item["course_name"], item["discounted_price"], 1, item["discounted_price"]])
    ws2.append([])
    ws2.append(["", "", "", "", "", "", "总计", round(total_all, 2)])

    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = 0
            for cell in col:
                try: max_len = max(max_len, len(str(cell.value or "")))
                except: pass
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 30)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = quote("教材订购明细_按宿舍.xlsx")
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"})



@router.get("/api/textbooks/orders/export-by-student-id")
def export_orders_by_student_id(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """按学号导出 Excel"""
    if not HAS_OPENPYXL:
        raise HTTPException(status_code=500, detail="服务器未安装 openpyxl")
    students = _get_student_dorm_data(db)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "教材订购明细(按学号)"
    header_font = openpyxl.styles.Font(bold=True, size=11, color="FFFFFF")
    header_fill = openpyxl.styles.PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center = openpyxl.styles.Alignment(horizontal="center", vertical="center")

    ws.append(["学号", "姓名", "宿舍", "教材名称", "课程名称", "单价", "折后价"])
    for cell in ws[1]: cell.font = header_font; cell.fill = header_fill; cell.alignment = center

    total_all = 0
    for sid in sorted(students.keys(), key=lambda x: students[x]["student_id"]):
        s = students[sid]
        for item in s["items"]:
            ws.append([s["student_id"], s["name"], s["dorm"], item["name"], item["course_name"], item["original_price"], item["discounted_price"]])
            total_all += item["discounted_price"]
        # 小计
        ws.append(["", "", "", "", "", "小计", round(sum(i["discounted_price"] for i in s["items"]), 2)])
        ws.cell(row=ws.max_row, column=6).font = openpyxl.styles.Font(bold=True)
        ws.cell(row=ws.max_row, column=7).font = openpyxl.styles.Font(bold=True)

    ws.append([])
    ws.append(["", "", "", "", "", "总计", round(total_all, 2)])
    ws.column_dimensions['A'].width = 14; ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 12; ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 16; ws.column_dimensions['F'].width = 10; ws.column_dimensions['G'].width = 10

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = quote("教材订购明细_按学号.xlsx")
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"})



@router.get("/api/textbooks/orders/export-by-dorm-word")
def export_orders_by_dorm_word(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """按宿舍导出 Word"""
    try:
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise HTTPException(status_code=500, detail="服务器未安装 python-docx")

    dorms = _get_dorm_grouped_data(db)
    students = _get_student_dorm_data(db)

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'SimSun'; style.font.size = Pt(10.5)

    title = doc.add_heading('数科2441班 教材订购明细（按宿舍）', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for dorm_name in sorted(dorms.keys()):
        data = dorms[dorm_name]
        doc.add_heading(f'{dorm_name} 宿舍', level=2)

        # 宿舍长查找
        dorm_leader = db.query(User).filter(User.dorm == dorm_name, User.is_dorm_leader == True).first()
        if dorm_leader:
            doc.add_paragraph(f'宿舍长：{dorm_leader.real_name}（{dorm_leader.student_id}）')

        # 订购汇总表
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Light List Accent 1'
        hdr = table.rows[0].cells
        hdr[0].text = '教材名称'; hdr[1].text = '课程名称'; hdr[2].text = '宿舍订购量'
        for hc in hdr: hc.paragraphs[0].runs[0].bold = True

        for tid, info in data["book_counts"].items():
            row = table.add_row().cells
            row[0].text = info["name"]; row[1].text = info["course_name"]; row[2].text = str(info["count"])

        # 订购学生名单
        dorm_students = [s for s in students.values() if s["dorm"] == dorm_name]
        if dorm_students:
            doc.add_paragraph(f'订购学生（共{len(dorm_students)}人）：')
            for s in sorted(dorm_students, key=lambda x: x["student_id"]):
                books = "、".join(f"{i['name']}（¥{i['discounted_price']}）" for i in s["items"])
                total = round(sum(i["discounted_price"] for i in s["items"]), 2)
                doc.add_paragraph(f'  {s["student_id"]} {s["name"]}：{books}  |  合计¥{total}')
        doc.add_paragraph()

    # 添加全部统计
    doc.add_page_break()
    doc.add_heading('全班级订购统计', level=2)
    total_books = sum(len(s["items"]) for s in students.values())
    total_amount = round(sum(sum(i["discounted_price"] for i in s["items"]) for s in students.values()), 2)
    doc.add_paragraph(f'订购总人数：{len(students)} 人')
    doc.add_paragraph(f'订购总本数：{total_books} 本')
    doc.add_paragraph(f'折后总金额：¥{total_amount}')

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    filename = quote("教材订购明细_按宿舍.docx")
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"})



@router.post("/api/textbooks/orders/send-dorm-emails")
def send_dorm_order_emails(db: Session = Depends(get_db, scope="function"), current_user: User = Depends(require_admin)):
    """一键发送宿舍订购明细到各宿舍长QQ邮箱"""
    students = _get_student_dorm_data(db)
    dorms = _get_dorm_grouped_data(db)

    sent = 0
    failed = 0
    for dorm_name in sorted(dorms.keys()):
        dorm_leader = db.query(User).filter(User.dorm == dorm_name, User.is_dorm_leader == True).first()
        if not dorm_leader or not dorm_leader.email or "@" not in dorm_leader.email:
            failed += 1
            continue

        data = dorms[dorm_name]
        dorm_students = [s for s in students.values() if s["dorm"] == dorm_name]

        lines = [f"【{dorm_name}宿舍·教材订购明细】\n"]
        lines.append(f"宿舍长：{dorm_leader.real_name}（{dorm_leader.student_id}）\n")
        lines.append("一、宿舍订购汇总：")
        total_count = 0
        for tid, info in data["book_counts"].items():
            lines.append(f"  {info['name']}（{info['course_name']}）：{info['count']} 本")
            total_count += info['count']
        lines.append(f"  合计：{total_count} 本\n")

        lines.append("二、各成员订购明细：")
        for s in sorted(dorm_students, key=lambda x: x["student_id"]):
            total_s = round(sum(i["discounted_price"] for i in s["items"]), 2)
            books = "、".join(f"{i['name']}（¥{i['discounted_price']}）" for i in s["items"])
            lines.append(f"  {s['student_id']} {s['name']}：{books}  |  合计¥{total_s}")

        total_dorm = round(sum(sum(i["discounted_price"] for i in s["items"]) for s in dorm_students), 2)
        lines.append(f"\n{len(dorm_students)} 人共 {total_count} 本，折后总额 ¥{total_dorm}")
        lines.append("\n请核对无误后按班级通知方式付款，如有疑问请联系班委。")

        try:
            result = send_emails_batch(
                [dorm_leader.email],
                f"【数科2441】{dorm_name}宿舍教材订购明细",
                "\n".join(lines),
                db=db, send_type="dorm_order", related_id=dorm_leader.id
            )
            if result["success"] and result["sent"] > 0:
                sent += 1
            else:
                failed += 1
        except Exception:
            failed += 1

    return {"message": f"已发送 {sent} 个宿舍" + (f"，{failed} 个失败" if failed > 0 else "")}

